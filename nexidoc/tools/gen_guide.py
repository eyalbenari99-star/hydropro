"""Nexi User Guide (.docx) — generated from the in-app Help centre (window.HNX_HELP_GUIDES in index.html).
Usage, from the repo root:
  python3 - <<'X'
import re;s=open('index.html',encoding='utf-8').read();i=s.find('<script id="hnx-help-guides-v1991">');st=s.find('>',i)+1;en=s.find('</script>',st);open('nexidoc/tools/guides.js','w').write(s[st:en])
X
  (cd nexidoc/tools && node dump-guides.js)   # writes nexidoc/tools/guides.json (dump-guides.js reads ./guides.js)
  python3 nexidoc/tools/gen_guide.py 21.29     # writes nexidoc/Nexi_User_Guide_v21.29.docx
Needs python-docx (pip install python-docx). The in-app guide is the source; edit it there, never here.
v23.08: every guide heading is followed by its "Added … — Updated …" line (the guide's added/updated stamps)."""
import json,re,sys,os,datetime
from docx import Document
from docx.shared import Pt, RGBColor
HERE=os.path.dirname(os.path.abspath(__file__))
ver=sys.argv[1] if len(sys.argv)>1 else 'dev'
G=json.load(open(os.path.join(HERE,'guides.json'),encoding='utf-8'))
def clean(t): return re.sub(r'<[^>]+>','',str(t if t is not None else ''))
MON=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
def stamp(s):
    """v23.08: {date:'2026-10-13',version:'22.88'} -> '13 Oct 2026 · v22.88' — the Help centre's words (no 🆕 on paper)."""
    s=s if isinstance(s,dict) else {}; m=re.match(r'^(\d{4})-(\d{2})-(\d{2})$',str(s.get('date','')))
    if not m or not 1<=int(m.group(2))<=12: return ''
    return '%d %s %s'%(int(m.group(3)),MON[int(m.group(2))-1],m.group(1))+(' · v'+str(s['version']) if s.get('version') else '')
def dated(g):
    a=stamp(g.get('added')); u=stamp(g.get('updated'))
    return ('Added '+a if a else '')+(' — ' if a and u else '')+('Updated '+u if u else '')
d=Document(); st=d.styles['Normal']; st.font.name='Calibri'; st.font.size=Pt(10.5)
d.add_heading('Nexi (HydroNexis-AI) — User Guide',0)
p=d.add_paragraph('ABA Pardes Agritech Corp. · Version '+ver+' · '+datetime.date.today().strftime('%-d %B %Y')); p.runs[0].font.color.rgb=RGBColor(0x55,0x66,0x60)
d.add_paragraph('Generated from the in-app Help centre (🆘). The in-app guide is the source; this document mirrors it screen by screen.')
d.add_heading('Contents',1)
for g in G: d.add_paragraph(clean(g.get('icon',''))+' '+clean(g.get('title','')),style='List Number')
def items(sec,key,style):
    for x in sec.get(key,[]) or []:
        if isinstance(x,list):
            para=d.add_paragraph(style=style); r=para.add_run(clean(x[0])+(': ' if len(x)>1 else '')); r.bold=len(x)>1
            if len(x)>1: para.add_run(clean(x[1]))
        else: d.add_paragraph(clean(x),style=style)
for g in G:
    d.add_page_break(); d.add_heading(clean(g.get('icon',''))+' '+clean(g.get('title','')),1)
    if dated(g): r=d.add_paragraph(dated(g)).runs[0]; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0x55,0x66,0x60)  # v23.08
    if g.get('who'): d.add_paragraph('For: '+clean(g['who'])).runs[0].italic=True
    if g.get('summary'): d.add_paragraph(clean(g['summary']))
    for s in g.get('sections',[]) or []:
        if s.get('h'): d.add_heading(clean(s['h']),2)
        items(s,'p',None) if False else None
        for x in s.get('p',[]) or []:
            if isinstance(x,list): items({'k':[x]},'k','List Bullet')
            else: d.add_paragraph(clean(x))
        items(s,'steps','List Number'); items(s,'items','List Bullet')
        t=s.get('table')
        if isinstance(t,dict) and t.get('rows'):
            head=t.get('head') or []; rows=t['rows']; cols=max(len(head),max(len(r) for r in rows))
            tb=d.add_table(rows=1 if head else 0,cols=cols); tb.style='Light Grid Accent 1'
            if head:
                for i,h in enumerate(head): c=tb.rows[0].cells[i]; c.text=clean(h); c.paragraphs[0].runs[0].bold=True
            for r in rows:
                cells=tb.add_row().cells
                for i,v in enumerate(r): cells[i].text=clean(v)
            d.add_paragraph('')
        for x in s.get('p2',[]) or []: d.add_paragraph(clean(x))
        n=s.get('note')
        if isinstance(n,dict):
            para=d.add_paragraph(); r=para.add_run((clean(n.get('k',''))+' — ') if n.get('k') else ''); r.bold=True; para.add_run(clean(n.get('t','')))
        elif n: d.add_paragraph(clean(n))
out=os.path.join(HERE,'..','Nexi_User_Guide_v'+ver+'.docx'); d.save(out); print('saved',out,len(G),'guides')
